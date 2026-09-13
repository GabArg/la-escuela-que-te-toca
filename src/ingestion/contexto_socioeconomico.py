"""Integra cuadros estáticos oficiales del Censo 2022 a nivel departamental."""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.ingestion.acceso_escolar import PROVINCES
from src.ingestion.educacion_departamental import normalize_text
from src.ingestion.ra_historico import load_bridge, load_georef

ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = ROOT / "data/raw/contexto_socioeconomico"
OUTPUT = ROOT / "data/processed/contexto_socioeconomico_2022.parquet"
COVERAGE = ROOT / "data/dictionaries/cobertura_contexto_socioeconomico.csv"
SOURCES = ROOT / "data/dictionaries/fuentes_contexto_socioeconomico.csv"
BRIDGE = ROOT / "data/dictionaries/puente_nombres_contexto_2022.csv"
BASE_URL = "https://www.indec.gob.ar/ftp/cuadros/poblacion/"
AGE_GROUPS = {"4_5": range(4, 6), "6_11": range(6, 12), "12_14": range(12, 15), "15_17": range(15, 18)}


def _number(value):
    """Convierte valores censales; `///` y otros signos permanecen ausentes."""
    return pd.to_numeric(value, errors="coerce")


def _province(path: Path) -> str:
    match = re.match(r"c2022_([^_]+)_", path.name)
    if not match or match.group(1) not in PROVINCES:
        raise ValueError(f"Jurisdicción no reconocida: {path.name}")
    return PROVINCES[match.group(1)]


def _department(title: str, province: str) -> str:
    match = re.search(r"(?:departamento|partido|comuna)\s+(.+?)\.\s*(?:Total|Población)", title, re.I)
    if not match:
        raise ValueError(f"No se pudo extraer territorio de: {title}")
    name = match.group(1).strip()
    return f"Comuna {name}" if normalize_text(province) == "ciudad autonoma de buenos aires" else name


def _sheets(path: Path):
    book = load_workbook(path, read_only=True, data_only=True)
    # Carátula, índice y total provincial preceden a las hojas departamentales.
    for sheet in book.worksheets[3:]:
        yield sheet


def _attach_territory(frame: pd.DataFrame) -> pd.DataFrame:
    data = frame.copy()
    data["provincia_key"] = data.provincia_fuente.map(normalize_text)
    data["departamento_key"] = data.departamento_fuente.map(normalize_text)
    geo = load_georef()
    cols = ["provincia_key", "departamento_key", "provincia_id", "provincia_nombre", "departamento_id", "departamento_nombre"]
    data = data.merge(geo[cols], on=["provincia_key", "departamento_key"], how="left", validate="many_to_one")
    bridge = load_bridge()
    bridge = bridge[(bridge.anio.astype("Int64") == 2022) & bridge.estado.eq("confirmado")]
    bridge = bridge[["provincia_key", "departamento_key", "departamento_id"]].rename(columns={"departamento_id": "bridge_id"})
    local = pd.read_csv(BRIDGE, dtype={"departamento_id": "string"})
    local = local[local.estado.eq("confirmado")].assign(
        provincia_key=lambda x: x.provincia.map(normalize_text),
        departamento_key=lambda x: x.nombre_censo.map(normalize_text),
    )[["provincia_key", "departamento_key", "departamento_id"]].rename(columns={"departamento_id": "bridge_id"})
    bridge = pd.concat([bridge, local], ignore_index=True).drop_duplicates(["provincia_key", "departamento_key"])
    data = data.merge(bridge, on=["provincia_key", "departamento_key"], how="left", validate="many_to_one")
    data["departamento_id"] = data.departamento_id.fillna(data.bridge_id)
    names = geo.set_index("departamento_id")
    data["provincia_id"] = data.provincia_id.fillna(data.departamento_id.map(names.provincia_id))
    data["provincia_nombre"] = data.provincia_nombre.fillna(data.departamento_id.map(names.provincia_nombre))
    data["departamento_nombre"] = data.departamento_nombre.fillna(data.departamento_id.map(names.departamento_nombre))
    return data


def load_household_context() -> pd.DataFrame:
    rows = []
    patterns = {"tecnologia": "*_hogares_c7_*.xlsx", "agua": "*_hogares_c2_*.xlsx", "saneamiento": "*_hogares_c3_*.xlsx"}
    for family, pattern in patterns.items():
        files = sorted(RAW_DIR.glob(pattern))
        if len(files) != 24:
            raise ValueError(f"{family}: se esperaban 24 archivos y se hallaron {len(files)}")
        for path in files:
            province = _province(path)
            for sheet in _sheets(path):
                title = str(sheet.cell(2, 1).value)
                department = _department(title, province)
                if family == "tecnologia":
                    total = _number(sheet.cell(5, 2).value)
                    computer = pd.Series([_number(sheet.cell(5, 4).value), _number(sheet.cell(5, 7).value)])
                    values = {"hogares_total": total, "hogares_internet": _number(sheet.cell(5, 3).value),
                              "hogares_computadora": computer.sum(min_count=2)}
                elif family == "agua":
                    values = {"hogares_agua_red_publica": _number(sheet.cell(6, 2).value)}
                else:
                    values = {"hogares_cloaca": _number(sheet.cell(6, 3).value)}
                rows.append({"provincia_fuente": province, "departamento_fuente": department, "familia": family, **values})
    data = pd.DataFrame(rows)
    data["provincia_key"] = data.provincia_fuente.map(normalize_text)
    data["departamento_key"] = data.departamento_fuente.map(normalize_text)
    identity = ["provincia_key", "departamento_key"]
    return data.groupby(identity, as_index=False, dropna=False).first().drop(columns=identity)


def load_attendance() -> pd.DataFrame:
    rows = []
    files = sorted(RAW_DIR.glob("*_educacion_c1_*.xlsx"))
    if len(files) != 24:
        raise ValueError(f"asistencia: se esperaban 24 archivos y se hallaron {len(files)}")
    for path in files:
        province = _province(path)
        for sheet in _sheets(path):
            department = _department(str(sheet.cell(2, 1).value), province)
            ages = {}
            for row in sheet.iter_rows(min_row=5, values_only=True):
                age = str(row[1]).strip() if row[1] is not None else ""
                if age.isdigit() and 4 <= int(age) <= 17:
                    ages[int(age)] = {key: pd.to_numeric(value, errors="coerce") for key, value in
                                      zip(["total", "asiste", "no_asiste_asistio", "nunca_asistio"], row[2:6])}
            for label, members in AGE_GROUPS.items():
                if not all(age in ages for age in members):
                    continue
                rows.append({"provincia_fuente": province, "departamento_fuente": department, "grupo_edad": label,
                             **{key: pd.Series([ages[age][key] for age in members]).sum(min_count=len(members))
                                for key in ages[next(iter(members))]}})
    return pd.DataFrame(rows)


def load_housing() -> pd.DataFrame:
    rows = []
    files = sorted(RAW_DIR.glob("*_vivienda_c3_*.xlsx"))
    if len(files) != 24:
        raise ValueError(f"vivienda: se esperaban 24 archivos y se hallaron {len(files)}")
    for path in files:
        province = _province(path); book = load_workbook(path, read_only=True, data_only=True)
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                code = f"{int(row[0]):05d}" if isinstance(row[0], (int, float)) else str(row[0]).strip()
                if not re.fullmatch(r"\d{5}", code):
                    continue
                precarious = pd.Series([_number(row[4]), _number(row[5])])
                rows.append({"provincia_fuente": province, "departamento_fuente": row[1],
                             "viviendas_particulares_ocupadas": _number(row[2]),
                             "viviendas_rancho_casilla": precarious.sum(min_count=2)})
    return pd.DataFrame(rows).drop_duplicates()


def build_context() -> pd.DataFrame:
    households = _attach_territory(load_household_context())
    attendance = _attach_territory(load_attendance())
    housing = _attach_territory(load_housing())
    att = attendance.pivot(index="departamento_id", columns="grupo_edad",
                           values=["total", "asiste", "no_asiste_asistio", "nunca_asistio"])
    att.columns = [f"{metric}_{group}" for metric, group in att.columns]
    att = att.reset_index()
    home = households.dropna(subset=["departamento_id"]).drop(
        columns=["familia", "provincia_fuente", "departamento_fuente", "provincia_key", "departamento_key", "bridge_id"]
    )
    identity = ["provincia_id", "provincia_nombre", "departamento_id", "departamento_nombre"]
    # Una misma unidad puede venir escrita de dos formas entre familias de cuadros;
    # se reúnen columnas complementarias, nunca se suman porcentajes ni conteos.
    home = home.groupby(identity, as_index=False, dropna=False).first()
    house = housing.dropna(subset=["departamento_id"]).groupby("departamento_id", as_index=False).first()
    house = house[["departamento_id", "viviendas_particulares_ocupadas", "viviendas_rancho_casilla"]]
    geo = load_georef()[["provincia_id", "provincia_nombre", "departamento_id", "departamento_nombre"]]
    out = geo.merge(home, on=["provincia_id", "provincia_nombre", "departamento_id", "departamento_nombre"], how="left", validate="one_to_one")
    out = out.merge(house, on="departamento_id", how="left", validate="one_to_one").merge(att, on="departamento_id", how="left", validate="one_to_one")
    out["porcentaje_hogares_internet"] = 100 * out.hogares_internet / out.hogares_total.where(out.hogares_total.gt(0))
    out["porcentaje_hogares_computadora"] = 100 * out.hogares_computadora / out.hogares_total.where(out.hogares_total.gt(0))
    out["porcentaje_hogares_agua_red_publica"] = 100 * out.hogares_agua_red_publica / out.hogares_total.where(out.hogares_total.gt(0))
    out["porcentaje_hogares_cloaca"] = 100 * out.hogares_cloaca / out.hogares_total.where(out.hogares_total.gt(0))
    out["porcentaje_viviendas_rancho_casilla"] = 100 * out.viviendas_rancho_casilla / out.viviendas_particulares_ocupadas.where(out.viviendas_particulares_ocupadas.gt(0))
    for group in AGE_GROUPS:
        out[f"porcentaje_asistencia_{group}"] = 100 * out[f"asiste_{group}"] / out[f"total_{group}"].where(out[f"total_{group}"].gt(0))
    out["anio_contexto"] = 2022
    return out


def write_source_inventory() -> None:
    family = {"hogares_c2": "agua", "hogares_c3": "saneamiento", "hogares_c7": "internet y computadora",
              "vivienda_c3": "tipo de vivienda", "educacion_c1": "asistencia escolar"}
    rows = []
    for path in sorted(RAW_DIR.glob("*.xlsx")):
        key = next(label for token, label in family.items() if token in path.name)
        rows.append({"fuente_id": path.stem, "organismo": "INDEC", "dataset": "Censo Nacional 2022", "variable": key,
                     "url": BASE_URL + path.name, "anio": 2022, "granularidad": "departamento/partido/comuna",
                     "universo": "según cuadro oficial y diccionario de variables", "formato": "XLSX", "fecha_consulta": "2026-09-12",
                     "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "observaciones": "Archivo original sin modificar"})
    rows.extend([
        {"fuente_id": "redatam_definiciones_2022", "organismo": "INDEC", "dataset": "Censo 2022 REDATAM",
         "variable": "NBI; hacinamiento; tecnologías; asistencia",
         "url": "https://redatam.indec.gob.ar/redarg/CENSOS/CPV2022/Docs/Redatam_Definiciones_de_la_base_de_datos.pdf",
         "anio": 2022, "granularidad": "consulta hasta radio censal", "universo": "hogares o personas según variable",
         "formato": "PDF/REDATAM", "fecha_consulta": "2026-09-12", "sha256": "",
         "observaciones": "Definiciones auditadas; NBI y hacinamiento no integrados"},
        {"fuente_id": "indec_nbi_2022", "organismo": "INDEC", "dataset": "Necesidades básicas insatisfechas",
         "variable": "NBI total y componentes", "url": "https://www.indec.gob.ar/indec/web/Nivel4-Tema-4-47-156",
         "anio": 2022, "granularidad": "publicación estática disponible a nivel provincial", "universo": "hogares",
         "formato": "HTML/XLSX", "fecha_consulta": "2026-09-12", "sha256": "",
         "observaciones": "Definición validada; sin cuadro estático departamental incorporable"},
    ])
    pd.DataFrame(rows).to_csv(SOURCES, index=False)


def write_outputs() -> pd.DataFrame:
    out = build_context(); OUTPUT.parent.mkdir(parents=True, exist_ok=True); out.to_parquet(OUTPUT, index=False)
    coverage = out[["provincia_id", "provincia_nombre", "departamento_id", "departamento_nombre"]].copy()
    coverage["nbi"] = "no incorporado"
    coverage["hacinamiento"] = "no incorporado"
    coverage["internet"] = out.porcentaje_hogares_internet.notna()
    coverage["computadora"] = out.porcentaje_hogares_computadora.notna()
    coverage["asistencia"] = out.porcentaje_asistencia_15_17.notna()
    coverage["vivienda_servicios"] = out.porcentaje_hogares_agua_red_publica.notna() & out.porcentaje_hogares_cloaca.notna() & out.porcentaje_viviendas_rancho_casilla.notna()
    coverage["observaciones"] = "Nulos preservados; NBI y hacinamiento excluidos por no contar con extracción estática departamental validada."
    coverage.to_csv(COVERAGE, index=False); write_source_inventory(); return out


if __name__ == "__main__":
    data = write_outputs(); print(data.notna().sum().to_string())
